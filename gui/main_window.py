import os
import tkinter as tk
from tkinter import filedialog, messagebox, ttk, simpledialog
from PIL import Image, ImageTk
import fitz
import re

# Importar el gestor de plantillas y el extractor de PDF
from templates.manager import TemplateManager
from processing.pdf_parser import PDFExtractor

class MainWindow:
    def __init__(self, root):
        self.root = root
        self.root.title("Extractor de Facturas")
        self._set_icon()  # Configurar el ícono/logo de la aplicación

        # Variables para almacenar la definición de la plantilla
        self.current_pdf = None
        self.current_page = 0
        self.selected_fields = []  # Lista de diccionarios para campos fijos
        self.tables = []           # Lista de definiciones de tablas
        self.mode = "field"        # "field" o "table"
        self.current_selection = None  # Área seleccionada en el canvas
        self.tk_img = None         # Referencia a la imagen en el canvas
        self.current_table = {}    # Diccionario temporal para la tabla en definición
        self.column_entries = []   # Para almacenar entradas en definición manual de columnas

        # Instanciar el gestor de plantillas (solo una vez)
        self.template_manager = TemplateManager()

        # Crear la interfaz y el menú
        self.create_widgets()
        self.create_menu()

    def _set_icon(self):
        """Configura el ícono de la aplicación."""
        try:
            # Primero, intenta usar un archivo ICO
            ico_path = os.path.join("assets", "IntervalorLogo.ico")
            if os.path.exists(ico_path):
                self.root.iconbitmap(ico_path)
            else:
                # Si el ICO no está disponible, intenta usar un PNG
                png_path = os.path.join("assets", "IntervalorLogo.png")
                if os.path.exists(png_path):
                    logo_img = ImageTk.PhotoImage(file=png_path)
                    self.root.iconphoto(False, logo_img)
                else:
                    print("Logo no encontrado en:", ico_path, "ni en:", png_path)
        except Exception as e:
            print("Error al cargar el logo:", e)

    def create_menu(self):
        """Crea la barra de menú, incluyendo la sección de Ayuda."""
        menubar = tk.Menu(self.root)
        help_menu = tk.Menu(menubar, tearoff=0)
        help_menu.add_command(label="Guía de Uso", command=self.show_help)
        help_menu.add_command(label="Acerca de", command=self.show_about)
        menubar.add_cascade(label="Ayuda", menu=help_menu)
        self.root.config(menu=menubar)

    def show_help(self):
        """Muestra una ventana emergente con la guía de uso de la aplicación."""
        help_text = (
            "Guía de Uso:\n\n"
            "1. Cargar PDF: Selecciona un archivo PDF que contenga la factura.\n"
            "2. Visualización: Se renderiza la primera página del PDF en el canvas.\n"
            "3. Seleccionar Área: Haz clic y arrastra sobre el canvas para definir el área de un campo o tabla.\n"
            "4. Agregar Campo: Ingresa el nombre del campo y haz clic en 'Agregar Campo' para guardar esa área.\n"
            "5. Nueva Tabla: Selecciona 'Nueva Tabla' para definir el área de una tabla y, luego, configura sus columnas.\n"
            "6. Guardar Template: Guarda la plantilla definida en un archivo JSON para reutilizarla.\n"
            "7. Procesar Documentos: Selecciona la plantilla y uno o varios PDFs para extraer los datos y exportarlos a Excel.\n\n"
            "Si tienes dudas adicionales, consulta la documentación o contacta al desarrollador."
        )
        messagebox.showinfo("Guía de Uso", help_text)

    def show_about(self):
        """Muestra información sobre la aplicación."""
        about_text = (
            "Extractor de Facturas\n"
            "Versión 1.0\n"
            "Desarrollado por Ariel Giamportone\n"
            "Esta aplicación permite extraer datos y tablas de facturas en PDF."
        )
        messagebox.showinfo("Acerca de", about_text)

    def create_widgets(self):
        # --- Barra superior ---
        self.top_frame = ttk.Frame(self.root)
        self.top_frame.pack(fill=tk.X, padx=5, pady=5)
        self._set_icon()

        # Agregar logo en la parte superior (opcional)
        logo_path = os.path.join("assets", "IntervalorLogo.png")
        if os.path.exists(logo_path):
            try:
                logo_img = Image.open(logo_path)
                logo_img = logo_img.resize((75, 35))  # Ajustar el tamaño según convenga
                self.tk_logo = ImageTk.PhotoImage(logo_img)
                logo_label = ttk.Label(self.top_frame, image=self.tk_logo)
                logo_label.pack(side=tk.LEFT, padx=5)
            except Exception as e:
                print("Error al cargar el logo en la interfaz:", e)

        self.btn_load = ttk.Button(self.top_frame, text="Cargar PDF", command=self.load_pdf)
        self.btn_load.pack(side=tk.LEFT, padx=5)
        self.field_name = ttk.Entry(self.top_frame, width=20)
        self.field_name.pack(side=tk.LEFT, padx=5)
        self.btn_add_field = ttk.Button(self.top_frame, text="Agregar Campo", command=self.add_field)
        self.btn_add_field.pack(side=tk.LEFT, padx=5)
        self.btn_edit_field = ttk.Button(self.top_frame, text="Editar Campo", command=self.edit_field)
        self.btn_edit_field.pack(side=tk.LEFT, padx=5)
        self.btn_new_table = ttk.Button(self.top_frame, text="Nueva Tabla", command=self.start_table_definition)
        self.btn_new_table.pack(side=tk.LEFT, padx=5)
        self.btn_define_columns = ttk.Button(self.top_frame, text="Definir Columnas", command=self.show_column_definition_dialog)
        self.btn_define_columns.pack(side=tk.LEFT, padx=5)
        self.btn_save_template = ttk.Button(self.top_frame, text="Guardar Template", command=self.save_template)
        self.btn_save_template.pack(side=tk.LEFT, padx=5)
        self.btn_process = ttk.Button(self.top_frame, text="Procesar Documentos", command=self.process_documents)
        self.btn_process.pack(side=tk.LEFT, padx=5)

        # --- Canvas y scrollbars ---
        self.canvas_frame = ttk.Frame(self.root)
        self.canvas_frame.pack(fill=tk.BOTH, expand=True)
        self.canvas = tk.Canvas(self.canvas_frame, bg="white")
        self.canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.scroll_x = ttk.Scrollbar(self.canvas_frame, orient=tk.HORIZONTAL, command=self.canvas.xview)
        self.scroll_y = ttk.Scrollbar(self.canvas_frame, orient=tk.VERTICAL, command=self.canvas.yview)
        self.scroll_x.pack(side=tk.BOTTOM, fill=tk.X)
        self.scroll_y.pack(side=tk.RIGHT, fill=tk.Y)
        self.canvas.configure(xscrollcommand=self.scroll_x.set, yscrollcommand=self.scroll_y.set)
        self.canvas.bind("<ButtonPress-1>", self.on_mouse_press)
        self.canvas.bind("<B1-Motion>", self.on_mouse_drag)
        self.canvas.bind("<ButtonRelease-1>", self.on_mouse_release)

        # --- Áreas de previsualización en un PanedWindow ---
        self.bottom_pane = ttk.PanedWindow(self.root, orient=tk.VERTICAL)
        self.bottom_pane.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)

        # Widget de previsualización de texto (por ejemplo, para mostrar mensajes o previews)
        self.preview_text = tk.Text(self.bottom_pane, height=3)
        # Listbox para listar campos y tablas definidos
        self.listbox = tk.Listbox(self.bottom_pane, height=3)

        # Agregar ambos widgets al PanedWindow con un peso (esto permite que sean redimensionables)
        self.bottom_pane.add(self.preview_text, weight=1)
        self.bottom_pane.add(self.listbox, weight=1)

    def load_pdf(self):
        """Carga un PDF y renderiza la primera página."""
        file_path = filedialog.askopenfilename(filetypes=[("PDF files", "*.pdf")])
        if file_path:
            try:
                self.current_pdf = fitz.open(file_path)
                self.current_page = 0
                self.render_pdf_page()
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo cargar el PDF:\n{e}")

    def render_pdf_page(self, page_num=0):
        """Renderiza la página indicada del PDF en el canvas."""
        if not self.current_pdf:
            return
        try:
            page = self.current_pdf.load_page(page_num)
            pix = page.get_pixmap()
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
            self.tk_img = ImageTk.PhotoImage(img)
            self.canvas.config(scrollregion=(0, 0, pix.width, pix.height))
            self.canvas.delete("all")
            self.canvas.create_image(0, 0, anchor=tk.NW, image=self.tk_img)
        except Exception as e:
            messagebox.showerror("Error", f"Error al renderizar la página:\n{e}")

    # --- Manejo de eventos del canvas ---
    def on_mouse_press(self, event):
        self.start_x = self.canvas.canvasx(event.x)
        self.start_y = self.canvas.canvasy(event.y)
        self.rect = self.canvas.create_rectangle(self.start_x, self.start_y, self.start_x, self.start_y, outline="red")

    def on_mouse_drag(self, event):
        cur_x = self.canvas.canvasx(event.x)
        cur_y = self.canvas.canvasy(event.y)
        self.canvas.coords(self.rect, self.start_x, self.start_y, cur_x, cur_y)

    def on_mouse_release(self, event):
        end_x = self.canvas.canvasx(event.x)
        end_y = self.canvas.canvasy(event.y)
        coords = (self.start_x, self.start_y, end_x, end_y)
        if self.mode == "field":
            self.current_selection = coords
            self.preview_text.insert(tk.END, f"Área de campo seleccionada: {coords}\n")
        elif self.mode == "table":
            self.current_table["coordinates"] = coords
            self.current_table["page"] = self.current_page
            self.preview_text.insert(tk.END, f"Área de tabla seleccionada: {coords}\n")
            self.auto_detect_columns()
            self.mode = "field"  # Regresa al modo campo

    # --- Definición y edición de campos fijos ---
    def add_field(self):
        """Agrega un campo fijo usando el nombre ingresado y el área seleccionada."""
        name = self.field_name.get().strip()
        if not name:
            messagebox.showerror("Error", "Debe ingresar un nombre para el campo")
            return
        if not self.current_selection:
            messagebox.showerror("Error", "Seleccione un área en el PDF para el campo")
            return
        field_def = {
            "name": name,
            "coordinates": self.current_selection,
            "page": self.current_page
        }
        self.selected_fields.append(field_def)
        self.listbox.insert(tk.END, f"Campo: {name} - {self.current_selection}")
        self.preview_text.insert(tk.END, f"Campo '{name}' agregado.\n")
        self.field_name.delete(0, tk.END)
        self.current_selection = None

    def edit_field(self):
        """Permite editar un campo previamente agregado."""
        selected_index = self.listbox.curselection()
        if not selected_index:
            messagebox.showerror("Error", "Seleccione un campo para editar")
            return
        index = selected_index[0]
        field = self.selected_fields.pop(index)
        self.listbox.delete(index)
        self.field_name.delete(0, tk.END)
        self.field_name.insert(0, field["name"])
        self.current_selection = field["coordinates"]
        self.preview_text.insert(tk.END, f"Edite el campo: {field['name']} (área: {field['coordinates']})\n")

    # --- Definición de tablas ---
    def start_table_definition(self):
        """Activa el modo de definición de tabla."""
        self.mode = "table"
        self.current_table = {"page": self.current_page}
        self.preview_text.insert(tk.END, "Seleccione el área de la tabla en el PDF.\n")

    def auto_detect_columns(self):
        """
        Tras definir el área de la tabla, se detectan columnas automáticamente.
        Se requiere implementar get_pdf_elements_in_area para extraer los elementos del PDF en esa zona.
        """
        if "coordinates" not in self.current_table:
            messagebox.showerror("Error", "Área de tabla no definida")
            return
        x0, y0, x1, y1 = self.current_table["coordinates"]
        pdf_elements = self.get_pdf_elements_in_area(x0, y0, x1, y1)
        detected_columns = self.detect_columns(pdf_elements)
        self.show_column_configurator(detected_columns)

    def get_pdf_elements_in_area(self, x0, y0, x1, y1):
        """
        Placeholder: implementar la lógica para extraer elementos (por ejemplo, bloques de texto)
        del PDF que se encuentren en el área dada.
        """
        return []

    def detect_columns(self, elements, tolerance=5):
        """
        Detecta columnas a partir de elementos con propiedades 'x0' y 'x1'.
        """
        x_positions = []
        for elem in elements:
            x_positions.append(elem.get('x0', 0))
            x_positions.append(elem.get('x1', 0))
        if not x_positions:
            return []
        rounded_x = [round(x / tolerance) * tolerance for x in x_positions]
        column_borders = sorted(list(set(rounded_x)))
        columns = []
        for i in range(len(column_borders) - 1):
            columns.append({
                'x0': column_borders[i],
                'x1': column_borders[i + 1],
                'name': f"Columna_{i + 1}"
            })
        return columns

    def show_column_configurator(self, columns):
        """
        Muestra una ventana para configurar las columnas detectadas.
        """
        self.column_configurator = tk.Toplevel(self.root)
        self.column_configurator.title("Configurar Columnas")
        self.config_canvas = tk.Canvas(self.column_configurator, width=800, height=600)
        self.config_canvas.pack()
        for col in columns:
            self.config_canvas.create_line(col['x0'], 0, col['x0'], 600, fill='red', dash=(2, 2))
        self.column_entries = []
        for i, col in enumerate(columns):
            frame = ttk.Frame(self.column_configurator)
            frame.pack(fill='x', padx=5, pady=2)
            ttk.Label(frame, text=f"Columna {i + 1}:").pack(side='left')
            entry = ttk.Entry(frame, width=20)
            entry.insert(0, col['name'])
            entry.pack(side='left')
            self.column_entries.append((col, entry))
        ttk.Button(self.column_configurator, text="Guardar Configuración", command=self.save_final_columns).pack(pady=10)

    def save_final_columns(self):
        """Guarda la configuración de columnas y pide el nombre de la tabla."""
        final_columns = []
        for col, entry in self.column_entries:
            final_columns.append({
                'x0': col['x0'],
                'x1': col['x1'],
                'name': entry.get().strip() or col['name']
            })
        self.current_table['columns'] = final_columns
        self.column_configurator.destroy()
        table_name = simpledialog.askstring("Nombre de Tabla", "Ingrese nombre para la tabla:")
        self.current_table['name'] = table_name or "Tabla_sin_nombre"
        self.tables.append(self.current_table)
        self.listbox.insert(tk.END, f"Tabla: {self.current_table['name']}")

    def show_column_definition_dialog(self):
        """
        Permite definir columnas de forma manual.
        """
        if "coordinates" not in self.current_table:
            messagebox.showerror("Error", "Primero defina el área de la tabla.")
            return
        dialog = tk.Toplevel(self.root)
        dialog.title("Definir Columnas de la Tabla")
        frame = ttk.Frame(dialog, padding=10)
        frame.pack()
        ttk.Label(frame, text="Nombre Columna").grid(row=0, column=0, padx=5, pady=5)
        self.column_entries = []
        for i in range(3):
            entry = ttk.Entry(frame, width=15)
            entry.grid(row=i + 1, column=0, padx=5, pady=2)
            self.column_entries.append(entry)
        add_row_btn = ttk.Button(frame, text="+ Añadir Fila", command=lambda: self.add_column_row(frame))
        add_row_btn.grid(row=4, column=0, padx=5, pady=5)
        save_btn = ttk.Button(frame, text="Guardar Columnas", command=lambda: self.save_table_columns(dialog))
        save_btn.grid(row=4, column=1, padx=5, pady=5)

    def add_column_row(self, frame):
        """Agrega una nueva fila de entrada para definir una columna."""
        next_row = len(self.column_entries) + 1
        entry = ttk.Entry(frame, width=15)
        entry.grid(row=next_row, column=0, padx=5, pady=2)
        self.column_entries.append(entry)

    def save_table_columns(self, dialog):
        """
        Guarda los nombres de columnas ingresados manualmente en la definición de la tabla.
        """
        final_columns = []
        for i, entry in enumerate(self.column_entries):
            col_name = entry.get().strip() or f"Columna_{i + 1}"
            final_columns.append({'name': col_name})
        self.current_table['columns'] = final_columns
        dialog.destroy()
        table_name = simpledialog.askstring("Nombre de Tabla", "Ingrese nombre para la tabla:")
        self.current_table['name'] = table_name or "Tabla_sin_nombre"
        self.tables.append(self.current_table)
        self.listbox.insert(tk.END, f"Tabla: {self.current_table['name']}")

    def save_template(self):
        """Guarda la plantilla en un archivo JSON."""
        if not self.selected_fields and not self.tables:
            messagebox.showwarning("Advertencia", "No hay campos ni tablas definidos")
            return
        template = {
            "fields": self.selected_fields,
            "tables": self.tables
        }
        file_path = filedialog.asksaveasfilename(
            defaultextension=".json",
            filetypes=[("JSON files", "*.json")]
        )
        if file_path:
            try:
                self.template_manager.save_template(template, file_path)
                messagebox.showinfo("Éxito", "Template guardado correctamente.")
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo guardar el template:\n{e}")

    def process_documents(self):
        """Procesa uno o varios PDFs usando un template guardado y genera un Excel."""
        template_path = filedialog.askopenfilename(filetypes=[("JSON files", "*.json")])
        if not template_path:
            return
        pdf_paths = filedialog.askopenfilenames(filetypes=[("PDF files", "*.pdf")])
        if not pdf_paths:
            return
        extractor = PDFExtractor(template_path)
        all_fields_data = []
        all_tables_data = []
        for pdf_path in pdf_paths:
            try:
                data = extractor.extract_from_pdf(pdf_path)
                field_data = {"Archivo": os.path.basename(pdf_path)}
                field_data.update(data.get("fields", {}))
                all_fields_data.append(field_data)
                for table_name, table_rows in data.get("tables", {}).items():
                    for row in table_rows:
                        row_data = {"Archivo": os.path.basename(pdf_path), "Tabla": table_name}
                        row_data.update(row)
                        all_tables_data.append(row_data)
            except Exception as e:
                messagebox.showwarning("Error", f"Error procesando {pdf_path}:\n{str(e)}")
        if all_fields_data or all_tables_data:
            self.export_to_excel(all_fields_data, all_tables_data)

    def export_to_excel(self, fields_data, tables_data):
        """Exporta la información extraída a un archivo Excel con dos hojas."""
        from openpyxl import Workbook
        from datetime import datetime
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        default_name = f"Procesado_{timestamp}.xlsx"
        output_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            initialfile=default_name
        )
        if output_path:
            try:
                wb = Workbook()
                ws_fields = wb.active
                ws_fields.title = "Campos"
                if fields_data:
                    headers = list(fields_data[0].keys())
                    ws_fields.append(headers)
                    for item in fields_data:
                        ws_fields.append([item.get(h, "") for h in headers])
                ws_tables = wb.create_sheet("Tablas")
                if tables_data:
                    headers = list(tables_data[0].keys())
                    ws_tables.append(headers)
                    for item in tables_data:
                        ws_tables.append([item.get(h, "") for h in headers])
                for ws in [ws_fields, ws_tables]:
                    for col in ws.columns:
                        max_length = 0
                        column = col[0].column_letter
                        for cell in col:
                            try:
                                cell_length = len(str(cell.value))
                                if cell_length > max_length:
                                    max_length = cell_length
                            except Exception:
                                pass
                        ws.column_dimensions[column].width = max_length + 2
                wb.save(output_path)
                messagebox.showinfo("Éxito", f"Documentos procesados.\nArchivo: {output_path}")
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo exportar a Excel:\n{e}")

def main():
    root = tk.Tk()
    app = MainWindow(root)
    root.mainloop()

if __name__ == "__main__":
    main()
