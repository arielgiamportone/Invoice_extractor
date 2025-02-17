import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from PIL import Image, ImageTk
import fitz  # PyMuPDF
import json
import re
import os

class App:
    def __init__(self, root):
        self.root = root
        self.root.title("Extractor de Facturas")
        self.current_pdf = None
        self.selected_areas = []
        self.tk_img = None  # Referencia para evitar garbage collection
        self.template = {"fields": [], "tables": []}
        self.current_table = None
        self.create_widgets()
        
    def create_widgets(self):
        # Frame superior
        self.top_frame = ttk.Frame(self.root)
        self.top_frame.pack(fill=tk.X, pady=5)
    
        self.btn_load = ttk.Button(self.top_frame, text="Cargar PDF", command=self.load_pdf)
        self.btn_load.pack(side=tk.LEFT, padx=5)
    
        self.field_name = ttk.Entry(self.top_frame)
        self.field_name.pack(side=tk.LEFT, padx=5)
    
        self.btn_add = ttk.Button(self.top_frame, text="Agregar campo", command=self.add_field)
        self.btn_add.pack(side=tk.LEFT, padx=5)
    
        self.btn_save = ttk.Button(self.top_frame, text="Guardar plantilla", command=self.save_template)
        self.btn_save.pack(side=tk.LEFT, padx=5)
    
        self.btn_new_table = ttk.Button(self.top_frame, text="Nueva Tabla", command=self.start_table_definition)
        self.btn_new_table.pack(side=tk.LEFT, padx=5)
    
        self.btn_define_columns = ttk.Button(self.top_frame, text="Definir Columnas", command=self.show_column_definition_dialog)
        self.btn_define_columns.pack(side=tk.LEFT, padx=5)
        
        self.btn_process = ttk.Button(self.top_frame, text="Procesar PDF", command=self.process_pdf)
        self.btn_process.pack(side=tk.LEFT, padx=5)
        
        # Frame para canvas y scrollbars
        self.canvas_frame = ttk.Frame(self.root)
        self.canvas_frame.pack(fill=tk.BOTH, expand=True)
        
        self.canvas = tk.Canvas(self.canvas_frame, bg="white")
        self.canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        
        self.scroll_x = ttk.Scrollbar(self.canvas_frame, orient=tk.HORIZONTAL, command=self.canvas.xview)
        self.scroll_y = ttk.Scrollbar(self.canvas_frame, orient=tk.VERTICAL, command=self.canvas.yview)
        self.scroll_x.pack(side=tk.BOTTOM, fill=tk.X)
        self.scroll_y.pack(side=tk.RIGHT, fill=tk.Y)
        
        self.canvas.configure(xscrollcommand=self.scroll_x.set, yscrollcommand=self.scroll_y.set)
        
        # Eventos para selección de áreas (campos)
        self.canvas.bind("<ButtonPress-1>", self.start_selection)
        self.canvas.bind("<B1-Motion>", self.draw_selection)
        self.canvas.bind("<ButtonRelease-1>", self.end_selection)
    
        # Área de previsualización de texto extraído
        self.preview_text = tk.Text(self.root, height=10)
        self.preview_text.pack(fill=tk.X, padx=5, pady=5)
    
        # ListBox para mostrar campos agregados
        self.list_frame = ttk.Frame(self.root)
        self.list_frame.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
    
        self.listbox = tk.Listbox(self.list_frame, height=6)
        self.listbox.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
    
        scrollbar = ttk.Scrollbar(self.list_frame, orient=tk.VERTICAL, command=self.listbox.yview)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.listbox.configure(yscrollcommand=scrollbar.set)
        
        # Label de estado
        self.status_label = ttk.Label(self.root, text="Listo")
        self.status_label.pack(fill=tk.X, padx=5, pady=2)
        
    def load_pdf(self):
        """Carga un archivo PDF y renderiza la primera página en el canvas."""
        file_path = filedialog.askopenfilename(filetypes=[("PDF files", "*.pdf")])
        if file_path:
            try:
                self.current_pdf = fitz.open(file_path)
                self.render_pdf_page()
                self.status_label.config(text="PDF cargado correctamente")
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo cargar el PDF:\n{e}")
    
    def render_pdf_page(self, page_num=0):
        """Renderiza la página indicada del PDF en el canvas."""
        if not self.current_pdf:
            return
        
        self.canvas.delete("all")  # Limpiar el canvas
        page = self.current_pdf.load_page(page_num)
        pix = page.get_pixmap()
        img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        
        self.tk_img = ImageTk.PhotoImage(img)
        self.canvas.config(scrollregion=(0, 0, pix.width, pix.height))
        self.canvas.create_image(0, 0, anchor=tk.NW, image=self.tk_img)
        
        # Opcional: volver a dibujar la plantilla (campos y tablas)
        self.draw_template_preview()
    
    def normalize_coords(self, x1, y1, x2, y2):
        """Asegura que las coordenadas estén ordenadas (x1,y1: superior izquierda, x2,y2: inferior derecha)."""
        return (min(x1, x2), min(y1, y2), max(x1, x2), max(y1, y2))
    
    # --- Funciones para la selección de áreas (campos) ---
    def start_selection(self, event):
        self.start_x = self.canvas.canvasx(event.x)
        self.start_y = self.canvas.canvasy(event.y)
        self.rect = None
    
    def draw_selection(self, event):
        cur_x = self.canvas.canvasx(event.x)
        cur_y = self.canvas.canvasy(event.y)
        
        if not self.rect:
            self.rect = self.canvas.create_rectangle(
                self.start_x, self.start_y, cur_x, cur_y, outline="red"
            )
        else:
            self.canvas.coords(self.rect, self.start_x, self.start_y, cur_x, cur_y)
    
    def end_selection(self, event):
        end_x = self.canvas.canvasx(event.x)
        end_y = self.canvas.canvasy(event.y)
        coords = self.normalize_coords(self.start_x, self.start_y, end_x, end_y)
        self.selected_areas.append({"coords": coords})
        self.status_label.config(text=f"Área seleccionada: {coords}")
    
    def add_field(self):
        """Asocia un nombre al área seleccionada y actualiza la previsualización."""
        name = self.field_name.get().strip()
        if not name:
            messagebox.showerror("Error", "Debe ingresar un nombre para el campo")
            return
        if not self.selected_areas:
            messagebox.showerror("Error", "No ha seleccionado ninguna área")
            return

        # Asigna el nombre al último área seleccionada
        self.selected_areas[-1]["name"] = name
        self.field_name.delete(0, tk.END)
        self.listbox.insert(tk.END, f"{name}: {self.selected_areas[-1]['coords']}")
    
        # Aquí se debe llamar a la función real de extracción de texto
        text = "Texto extraído (simulado)"
        self.preview_text.insert(tk.END, f"{name}: {text}\n")
    
    def draw_template_preview(self):
        """Dibuja en el canvas las áreas definidas para campos y tablas."""
        # Dibujar campos
        for field in self.template.get("fields", []):
            x1, y1, x2, y2 = field["coordinates"]
            self.canvas.create_rectangle(x1, y1, x2, y2, outline="green")
        # Dibujar tablas
        for table in self.template.get("tables", []):
            x1, y1, x2, y2 = table["coordinates"]
            self.canvas.create_rectangle(x1, y1, x2, y2, outline="blue", width=2)
            self.canvas.create_text(x1+10, y1+10, text=table["name"], fill="blue", anchor="nw")
    
    def save_template(self):
        """Guarda la plantilla actual en un archivo JSON."""
        if not self.selected_areas and not hasattr(self, 'tables'):
            messagebox.showwarning("Advertencia", "No hay campos o tablas definidos")
            return
    
        valid_fields = [area for area in self.selected_areas if 'name' in area]
    
        template = {
            "page_number": 0,
            "fields": [{"name": area["name"], "coordinates": list(area["coords"])} for area in valid_fields],
            "tables": getattr(self, 'tables', [])
        }
    
        file_path = filedialog.asksaveasfilename(
            defaultextension=".json",
            filetypes=[("JSON files", "*.json")]
        )
    
        if file_path:
            try:
                with open(file_path, "w") as f:
                    json.dump(template, f, indent=4)
                messagebox.showinfo("Éxito", "Plantilla guardada correctamente")
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo guardar la plantilla:\n{e}")
    
    def process_pdf(self):
        """Procesa los PDFs seleccionados utilizando la plantilla cargada."""
        # Importar la lógica de extracción
        from processing.extractor import PDFExtractor
        
        # Seleccionar plantilla
        template_path = filedialog.askopenfilename(
            filetypes=[("Plantillas JSON", "*.json")]
        )
        if not template_path:
            return
        
        # Seleccionar uno o varios PDFs
        pdf_paths = filedialog.askopenfilenames(
            title="Seleccionar archivos PDF",
            filetypes=[("Archivos PDF", "*.pdf")]
        )
        if not pdf_paths:
            return
        
        extractor = PDFExtractor(template_path)
        all_data = []
        
        for pdf_path in pdf_paths:
            try:
                data = extractor.extract_from_pdf(pdf_path)
                data["Archivo"] = os.path.basename(pdf_path)
                all_data.append(data)
            except Exception as e:
                messagebox.showwarning("Advertencia", f"Error en {pdf_path}:\n{str(e)}")
        
        if all_data:
            self.export_to_excel(all_data)
    
    def export_to_excel(self, data):
        """Exporta la información extraída a un archivo Excel."""
        from openpyxl import Workbook
        from datetime import datetime
    
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        default_name = f"Facturas_Procesadas_{timestamp}.xlsx"
    
        output_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel Files", "*.xlsx")],
            initialfile=default_name
        )
    
        if output_path:
            try:
                wb = Workbook()
                ws = wb.active
            
                headers = list(data[0].keys())
                ws.append(headers)
            
                for item in data:
                    ws.append([item.get(h, "") for h in headers])
            
                # Autoajustar columnas
                for col in ws.columns:
                    max_length = 0
                    column = col[0].column_letter
                    for cell in col:
                        try:
                            cell_length = len(str(cell.value))
                            if cell_length > max_length:
                                max_length = cell_length
                        except Exception:
                            pass
                    adjusted_width = (max_length + 2)
                    ws.column_dimensions[column].width = adjusted_width
                
                wb.save(output_path)
                messagebox.showinfo("Éxito", f"{len(data)} documentos procesados:\n{output_path}")
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo exportar a Excel:\n{e}")
    
    # --- Funciones para la definición de tablas ---
    def start_table_definition(self):
        """Activa el modo para definir una nueva tabla."""
        self.current_table = {
            "name": "",
            "area": None,
            "columns": [],
            "start_marker": "",
            "end_marker": ""
        }
        # Por ejemplo, se puede solicitar el nombre de la tabla mediante un Entry o un diálogo
        self.table_name_entry = ttk.Entry(self.root)
        self.table_name_entry.pack(padx=5, pady=2)
        messagebox.showinfo("Tabla", "Modo de definición de tabla activado. Seleccione el área en el PDF.")
        self.canvas.bind("<Button-1>", self.select_table_area_start)
        self.status_label.config(text="Dibuje el área de la tabla en el PDF")
    
    def select_table_area_start(self, event):
        self.table_start_x = self.canvas.canvasx(event.x)
        self.table_start_y = self.canvas.canvasy(event.y)
    
        self.table_rect = self.canvas.create_rectangle(
            self.table_start_x, 
            self.table_start_y,
            self.table_start_x,
            self.table_start_y,
            outline="#2A9DF4",
            width=2
        )
        self.canvas.bind("<B1-Motion>", self.select_table_area_update)
        self.canvas.bind("<ButtonRelease-1>", self.select_table_area_end)
    
    def select_table_area_update(self, event):
        current_x = self.canvas.canvasx(event.x)
        current_y = self.canvas.canvasy(event.y)
        self.canvas.coords(
            self.table_rect,
            self.table_start_x,
            self.table_start_y,
            current_x,
            current_y
        )
    
    def select_table_area_end(self, event):
        end_x = self.canvas.canvasx(event.x)
        end_y = self.canvas.canvasy(event.y)
        self.current_table["area"] = self.normalize_coords(self.table_start_x, self.table_start_y, end_x, end_y)
        self.canvas.unbind("<B1-Motion>")
        self.canvas.unbind("<ButtonRelease-1>")
        self.status_label.config(text="Área de tabla definida")
    
        self.current_table["name"] = self.table_name_entry.get().strip()
        if not re.match(r"^[\w\s]+$", self.current_table["name"]):
            messagebox.showerror("Error", "Nombre de tabla no válido")
            return
        # Aquí podrías agregar la tabla a la plantilla o a una lista de tablas
        if not hasattr(self, 'tables'):
            self.tables = []
        # Suponiendo que el área de la tabla es la que se usará en la exportación
        table_def = {
            "name": self.current_table["name"],
            "coordinates": list(self.current_table["area"]),
            "columns": self.current_table["columns"]
        }
        self.tables.append(table_def)
        # Actualizar previsualización si es necesario
        self.draw_template_preview()
    
    def show_column_definition_dialog(self):
        """Abre un diálogo para que el usuario defina las columnas de la tabla."""
        column_dialog = tk.Toplevel(self.root)
        column_dialog.title("Definir Columnas")
        # Aquí se puede implementar la lógica para agregar, editar y eliminar columnas
        # Por ejemplo, usando Listbox y Entry para cada columna
        ttk.Label(column_dialog, text="Definir columnas (en desarrollo)").pack(padx=10, pady=10)
    
    def load_template(self):
        # TODO: Implementar lógica para cargar una plantilla desde un archivo JSON
        pass

def main_gui():
    root = tk.Tk()
    App(root)
    root.mainloop()

if __name__ == "__main__":
    main_gui()
